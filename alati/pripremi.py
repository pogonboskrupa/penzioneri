# -*- coding: utf-8 -*-
"""Čisti sve spiskove i priprema tabele za Google Sheets.

VAŽNO: svaka grupa korisnika (spisak) se vodi ODVOJENO - ima svoj odobreni
fond kubika, svoj list PODACI_* i ULICE_*:

    CIJEPANO         podaci/DRVA_PENZ_2025.xls            (KOL u prostornim metrima)
    U DUGOM          podaci/PENZ_BOS_KRUPA_U_DUGOM.xls    (KOL već u m3)
    RVI              podaci/RVI_BOSANSKA_KRUPA_2026.xls   (kolona "količina" = odobreno)
    PORODICE ŠEHIDA  podaci/PORODICE_SEHIDA_2025.xls      (KOL već u m3)
    SINDIKAT         podaci/SINDIKAT.csv                  (ručno pripremljen iz .doc)

Pokretanje:
    python3 alati/pripremi.py

Rezultat (folder izlaz/):
    PENZIONERI_DRVA.xlsx  - PODACI_<grupa>, ULICE_<grupa> za svaku grupu,
                            plus MJESTA i ULICE_GEO (zajednički, za kartu)
    *.csv                 - isti listovi za uvoz u Google Sheets
"""
import os
import sys
from collections import Counter, defaultdict

import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from normalizacija import (  # noqa: E402
    broj,
    kolicina_u_m3,
    naziv_ulice_za_prikaz,
    normalizuj_mjesto,
    normalizuj_ulicu,
    spoji_slicne,
    uljepsaj_prikaz,
)

KORIJEN = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# "kod" = ASCII sufiks za nazive listova/kolona (bez razmaka i dijakritike);
# "naziv" = čitljiv naziv koji ide u kolonu Vrsta i u izbore na stranici;
# "nacin": prm (KOL u prostornim metrima) | m3 (KOL već u m3) |
#          rvi (kolona "količina" = odobreno, bez posebne KOL kolone) |
#          sindikat (poseban CSV, vidi napravi_sindikat_csv.py)
IZVORI = [
    {"kod": "CIJEPANO", "naziv": "CIJEPANO", "nacin": "prm",
     "putanja": "podaci/DRVA_PENZ_2025.xls", "kategorija": "PENZIONER"},
    {"kod": "U_DUGOM", "naziv": "U DUGOM", "nacin": "m3",
     "putanja": "podaci/PENZ_BOS_KRUPA_U_DUGOM.xls", "kategorija": "PENZIONER"},
    {"kod": "RVI", "naziv": "RVI", "nacin": "rvi",
     "putanja": "podaci/RVI_BOSANSKA_KRUPA_2026.xls", "kategorija": "RVI"},
    {"kod": "PORODICE_SEHIDA", "naziv": "PORODICE ŠEHIDA", "nacin": "m3",
     "putanja": "podaci/PORODICE_SEHIDA_2025.xls", "kategorija": "PORODICA ŠEHIDA"},
    {"kod": "SINDIKAT", "naziv": "SINDIKAT", "nacin": "sindikat",
     "putanja": "podaci/SINDIKAT.csv", "kategorija": "SINDIKAT"},
]

KOLONE = ["rb", "maticni", "prezime", "ime", "mjesto", "ulica", "telefon",
          "kol", "datum", "otpremnica", "kolicina", "oznaka"]

REDOSLIJED = ["Vrsta", "Redni broj", "Matični broj", "Prezime", "Ime", "Kategorija",
              "Mjesto",
              "Ulica", "Telefon", "Odobreno m3", "Isporučeno m3", "Preostalo m3",
              "Status", "Datum isporuke", "Otpremnica", "Napomena",
              "Adresa ključ", "Original mjesto", "Original ulica"]


def redni_broj(v):
    """Svede 'Redni broj' na pravi cijeli broj (ne tekst) da bi sortiranje i
    filtriranje u Google Sheetsu bilo brojčano (1,2,3...), a ne
    leksikografsko (1,10,11,2...). Izvorni zapisi znaju imati tačku na
    kraju ("1.") ili biti već float (110.0)."""
    s = str(v).strip().rstrip(".")
    try:
        return int(float(s))
    except ValueError:
        return v


def _red(vrsta, kategorija, rb, maticni, prezime, ime, mjesto_sirovo, ulica_sirovo,
         telefon, odobreno, isporuceno, datum, otpremnica, napomena):
    odobreno = round(odobreno or 0.0, 2)
    isporuceno = round(isporuceno or 0.0, 2)
    preostalo = max(round(odobreno - isporuceno, 2), 0.0)
    return {
        "Vrsta": vrsta,
        "Redni broj": redni_broj(rb),
        "Matični broj": maticni,
        "Prezime": prezime,
        "Ime": ime,
        "Kategorija": kategorija,
        "Mjesto": normalizuj_mjesto(mjesto_sirovo) if mjesto_sirovo else "Bosanska Krupa",
        "_ulica_kljuc": normalizuj_ulicu(ulica_sirovo) or "(BEZ ULICE)",
        "_ulica_original": str(ulica_sirovo).strip() if ulica_sirovo else "",
        "Telefon": telefon or "",
        "Odobreno m3": odobreno,
        "Isporučeno m3": isporuceno,
        "Preostalo m3": preostalo,
        "Status": ("ISPORUČENO" if preostalo <= 0.001 and isporuceno > 0
                   else "DJELIMIČNO" if isporuceno > 0 else "ZA ISPORUKU"),
        "Datum isporuke": datum or "",
        "Otpremnica": otpremnica or "",
        "Napomena": napomena or "",
        "Original mjesto": str(mjesto_sirovo).strip() if mjesto_sirovo else "",
        "Original ulica": str(ulica_sirovo).strip() if ulica_sirovo else "",
    }


def ucitaj_xls(izvor):
    """CIJEPANO / U DUGOM / PORODICE ŠEHIDA (nacin prm ili m3) i RVI (nacin rvi)."""
    sirovo = pd.read_excel(os.path.join(KORIJEN, izvor["putanja"]),
                           sheet_name="Naknadni spisak", header=None)
    d = sirovo.iloc[2:, :12].copy()
    d.columns = KOLONE

    redovi = []
    for _, r in d.iterrows():
        prezime = str(r.prezime).strip() if pd.notna(r.prezime) else ""
        ime = str(r.ime).strip() if pd.notna(r.ime) else ""
        if not prezime and not ime:
            continue
        datum = str(r.datum).strip() if pd.notna(r.datum) else ""
        otpremnica = str(r.otpremnica).strip() if pd.notna(r.otpremnica) else ""

        if izvor["nacin"] == "rvi":
            # nema posebne "odobreno" kolone - "količina" nosi odobreni iznos,
            # a isporuka se prepoznaje po datumu/otpremnici (isti iznos)
            odobreno = broj(r.kolicina) or 0.0
            isporuceno = odobreno if (datum or otpremnica) else 0.0
        else:
            odobreno = kolicina_u_m3(r.kol, izvor["nacin"] == "m3") or 0.0
            isporuceno = broj(r.kolicina)
            if isporuceno is None and (datum or otpremnica):
                isporuceno = odobreno
            isporuceno = isporuceno or 0.0

        redovi.append(_red(
            izvor["naziv"], izvor["kategorija"],
            str(r.rb).strip().replace(".0", "") if pd.notna(r.rb) else "",
            str(r.maticni).strip().replace(".0", "") if pd.notna(r.maticni) else "",
            prezime, ime, r.mjesto if pd.notna(r.mjesto) else "",
            r.ulica if pd.notna(r.ulica) else "",
            str(r.telefon).strip() if pd.notna(r.telefon) else "",
            odobreno, isporuceno, datum, otpremnica,
            str(r.oznaka).strip() if pd.notna(r.oznaka) else ""))
    return pd.DataFrame(redovi)


def ucitaj_sindikat(izvor):
    """SINDIKAT - ručno pripremljen CSV (vidi napravi_sindikat_csv.py)."""
    d = pd.read_csv(os.path.join(KORIJEN, izvor["putanja"]))
    redovi = []
    for _, r in d.iterrows():
        ime_puno = str(r.prezime_ime).strip()
        dijelovi = ime_puno.split(" ", 1)
        prezime = dijelovi[0] if dijelovi else ime_puno
        ime = dijelovi[1] if len(dijelovi) > 1 else ""
        redovi.append(_red(
            izvor["naziv"], izvor["kategorija"], str(r.rb), "",
            prezime, ime, r.mjesto, "",
            str(r.telefon).strip() if pd.notna(r.telefon) else "",
            r.odobreno_m3, r.isporuceno_m3,
            str(r.datum).strip() if pd.notna(r.datum) else "",
            str(r.otpremnica).strip() if pd.notna(r.otpremnica) else "",
            str(r.napomena).strip() if pd.notna(r.napomena) else ""))
    return pd.DataFrame(redovi)


def ucitaj(izvor):
    if izvor["nacin"] == "sindikat":
        return ucitaj_sindikat(izvor)
    return ucitaj_xls(izvor)


def ujednaci_ulice(podaci):
    """Spaja različite zapise iste ulice, ali samo unutar istog mjesta."""
    po_mjestu = defaultdict(Counter)
    for m, k in zip(podaci["Mjesto"], podaci["_ulica_kljuc"]):
        po_mjestu[m][k] += 1
    mapa = spoji_slicne(po_mjestu)

    # za svaki glavni ključ biramo najljepši originalni zapis (s dijakritikom)
    kandidati = defaultdict(Counter)
    for m, k, o in zip(podaci["Mjesto"], podaci["_ulica_kljuc"],
                       podaci["_ulica_original"]):
        glavni = mapa[(m, k)]
        prikaz = uljepsaj_prikaz(naziv_ulice_za_prikaz(o)) if o else "(bez ulice)"
        if prikaz:
            kandidati[(m, glavni)][prikaz] += 1

    def najbolji(par):
        c = kandidati.get(par)
        if not c:
            return "(bez ulice)"
        # prednost zapisu s dijakritikom, pa češćem, pa dužem
        return sorted(c, key=lambda p: (-any(z in p for z in "čćžšđČĆŽŠĐ"),
                                        -c[p], -len(p), p))[0]

    glavni = [mapa[(m, k)] for m, k in zip(podaci["Mjesto"], podaci["_ulica_kljuc"])]
    podaci["Ulica"] = [najbolji((m, g)) for m, g in zip(podaci["Mjesto"], glavni)]
    podaci["Adresa ključ"] = [f"{m}|{g}" for m, g in zip(podaci["Mjesto"], glavni)]
    return podaci


def sazetak_ulica(podaci):
    u = (podaci.groupby(["Mjesto", "Ulica", "Adresa ključ"], as_index=False)
         .agg(**{"Broj korisnika": ("Prezime", "count"),
                 "Odobreno m3": ("Odobreno m3", "sum"),
                 "Isporučeno m3": ("Isporučeno m3", "sum"),
                 "Preostalo m3": ("Preostalo m3", "sum")}))
    gotovi = (podaci[podaci["Status"] == "ISPORUČENO"]
              .groupby("Adresa ključ").size())
    u["Isporučeno korisnika"] = [int(gotovi.get(k, 0)) for k in u["Adresa ključ"]]
    u["Za isporuku korisnika"] = u["Broj korisnika"] - u["Isporučeno korisnika"]
    u["Status ulice"] = ["ZAVRŠENO" if p <= 0.001 else
                         ("U TOKU" if i > 0.001 else "NIJE POČETO")
                         for p, i in zip(u["Preostalo m3"], u["Isporučeno m3"])]
    for k in ("Odobreno m3", "Isporučeno m3", "Preostalo m3"):
        u[k] = u[k].round(2)
    return u.sort_values(["Preostalo m3", "Mjesto", "Ulica"],
                         ascending=[False, True, True]).reset_index(drop=True)


BOJA_ISPORUCENO = "FFD9EAD3"    # svijetlo zelena
BOJA_NEISPORUCENO = "FFF4CCCC"  # svijetlo crvena


def oboji_listove(writer, listovi):
    """Cijeli red korisnika/ulice u zelenoj ili crvenoj nijansi."""
    from openpyxl.styles import PatternFill

    zelena = PatternFill("solid", fgColor=BOJA_ISPORUCENO)
    crvena = PatternFill("solid", fgColor=BOJA_NEISPORUCENO)
    for naziv, df in listovi.items():
        if "Preostalo m3" not in df.columns:
            continue
        list_ = writer.sheets[naziv]
        for i, preostalo in enumerate(df["Preostalo m3"], start=2):
            ispuna = zelena if float(preostalo) <= 0.001 else crvena
            for kolona in range(1, len(df.columns) + 1):
                list_.cell(row=i, column=kolona).fill = ispuna


def main():
    dijelovi = {izvor["naziv"]: ucitaj(izvor) for izvor in IZVORI}
    sve = pd.concat(dijelovi.values(), ignore_index=True)
    sve = ujednaci_ulice(sve)
    sve = sve.sort_values(["Vrsta", "Mjesto", "Ulica", "Prezime", "Ime"],
                          kind="stable").reset_index(drop=True)

    po_vrsti = {v: sve[sve["Vrsta"] == v][REDOSLIJED].reset_index(drop=True)
                for v in dijelovi}
    ulice = {v: sazetak_ulica(d) for v, d in po_vrsti.items()}

    mjesta = (sve.groupby(["Vrsta", "Mjesto"], as_index=False)
              .agg(**{"Broj korisnika": ("Prezime", "count"),
                      "Odobreno m3": ("Odobreno m3", "sum"),
                      "Isporučeno m3": ("Isporučeno m3", "sum"),
                      "Preostalo m3": ("Preostalo m3", "sum")})
              .sort_values(["Vrsta", "Preostalo m3"], ascending=[True, False]))
    for k in ("Odobreno m3", "Isporučeno m3", "Preostalo m3"):
        mjesta[k] = mjesta[k].round(2)

    # jedan red po ulici (sve grupe zajedno) - osnova za kartu
    geo = (sve.groupby(["Mjesto", "Ulica", "Adresa ključ"], as_index=False)
           .size().drop(columns="size"))
    kolone_preostalo = []
    for izvor in IZVORI:
        v = izvor["naziv"]
        s = ulice[v].set_index("Adresa ključ")
        kolona = f"Preostalo {v} m3"
        geo[kolona] = [float(s["Preostalo m3"].get(k, 0.0)) for k in geo["Adresa ključ"]]
        geo[f"Isporučeno {v} m3"] = [float(s["Isporučeno m3"].get(k, 0.0))
                                     for k in geo["Adresa ključ"]]
        kolone_preostalo.append(kolona)
    geo["Preostalo ukupno m3"] = geo[kolone_preostalo].sum(axis=1).round(2)
    geo["Adresa za kartu"] = [
        (f"{u}, {m}, Bosna i Hercegovina" if u != "(bez ulice)"
         else f"{m}, Bosna i Hercegovina")
        for u, m in zip(geo["Ulica"], geo["Mjesto"])]
    geo["Lat"] = ""
    geo["Lng"] = ""
    geo = geo.sort_values("Preostalo ukupno m3", ascending=False).reset_index(drop=True)

    izlaz = os.path.join(KORIJEN, "izlaz")
    os.makedirs(izlaz, exist_ok=True)
    listovi = {}
    for izvor in IZVORI:
        listovi[f"PODACI_{izvor['kod']}"] = po_vrsti[izvor["naziv"]]
        listovi[f"ULICE_{izvor['kod']}"] = ulice[izvor["naziv"]]
    listovi["MJESTA"] = mjesta
    listovi["ULICE_GEO"] = geo

    with pd.ExcelWriter(os.path.join(izlaz, "PENZIONERI_DRVA.xlsx"),
                        engine="openpyxl") as w:
        for naziv, df in listovi.items():
            df.to_excel(w, sheet_name=naziv, index=False)
        oboji_listove(w, listovi)
    for naziv, df in listovi.items():
        df.to_csv(os.path.join(izlaz, f"{naziv}.csv"), index=False)

    for izvor in IZVORI:
        d = po_vrsti[izvor["naziv"]]
        print(f"[{izvor['naziv']}] korisnika {len(d)}, ulica {len(ulice[izvor['naziv']])}, "
              f"odobreno {d['Odobreno m3'].sum():.2f} m3, "
              f"isporučeno {d['Isporučeno m3'].sum():.2f} m3, "
              f"preostalo {d['Preostalo m3'].sum():.2f} m3")
    print(mjesta.to_string(index=False))


if __name__ == "__main__":
    main()
